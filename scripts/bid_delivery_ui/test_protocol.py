#!/usr/bin/env python3
"""Lifecycle tests for the host-neutral delivery manifest protocol."""

from __future__ import annotations

import json
import os
import sys
import tempfile
import unittest
import zipfile
from copy import deepcopy
from pathlib import Path


SCRIPT_DIR = Path(__file__).resolve().parent
SCRIPTS_DIR = SCRIPT_DIR.parent
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from bid_confirm_ui import server as confirm_ui
from bid_delivery_ui import export_image_plan, export_word, protocol


class DeliveryProtocolTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.project_dir = Path(self.temp.name)
        self.data_dir = self.project_dir / confirm_ui.DATA_DIR_NAME
        self.data_dir.mkdir()
        self._write_confirmed_authorization()

    def tearDown(self) -> None:
        self.temp.cleanup()

    def _write(self, name: str, value: dict) -> None:
        confirm_ui.atomic_write_json(self.data_dir / name, value)

    def _receipt(self, name: str, value: dict) -> dict:
        value["confirmation_sha256"] = confirm_ui.sha256_data(value)
        self._write(name, value)
        return value

    def _write_confirmed_authorization(self) -> None:
        project_id = "delivery-protocol-test"
        stage1 = {
            "schema_version": 1, "stage": "stage1", "project_id": project_id,
            "project": {"project_name": "测试标书"}, "scoring": {},
            "formatting": {"target_pages": 20}, "boundaries": {},
        }
        self._write(confirm_ui.STAGE1_INPUT, stage1)
        stage1_receipt = self._receipt(confirm_ui.STAGE1_RECEIPT, {
            "schema_version": 1, "stage": "stage1", "status": "confirmed", "project_id": project_id,
            "source_sha256": confirm_ui.sha256_data(stage1), "data": stage1,
            "confirmed_at": confirm_ui.utc_now(),
        })
        chapters = [
            {"id": "chapter-1", "number": "1", "title": "第1章", "level": 1, "order": 1, "pages": 10, "score_refs": [], "requirement_refs": [], "allow_deeper": False, "children": []},
            {"id": "chapter-2", "number": "2", "title": "第2章", "level": 1, "order": 2, "pages": 10, "score_refs": [], "requirement_refs": [], "allow_deeper": False, "children": []},
        ]
        stage2 = {
            "schema_version": 1, "stage": "stage2", "project_id": project_id,
            "stage1_confirmation_sha256": stage1_receipt["confirmation_sha256"], "target_pages": 20,
            "coverage": {"total": 0, "mapped": 0, "unmapped": []}, "chapters": chapters,
        }
        self._write(confirm_ui.STAGE2_INPUT, stage2)
        stage2_receipt = self._receipt(confirm_ui.STAGE2_RECEIPT, {
            "schema_version": 1, "stage": "stage2", "status": "confirmed", "project_id": project_id,
            "stage1_confirmation_sha256": stage1_receipt["confirmation_sha256"],
            "source_sha256": confirm_ui.sha256_data(stage2),
            "data": {"chapters": chapters, "coverage": stage2["coverage"], "planned_pages": 20},
            "confirmed_at": confirm_ui.utc_now(),
        })
        image = {
            "id": "image-1", "figure_no": "图1-1", "order": 1, "chapter_id": "chapter-1",
            "chapter_number": "1", "chapter_title": "第1章",
            "position": {"outline_node_id": "chapter-1", "outline_number": "1", "outline_title": "第1章", "placement_note": "章导语后"},
            "name": "总览图", "type": "章首总览图", "purpose": "概括方案", "core_nodes": ["目标"],
            "composition": "分层结构", "orientation": "landscape", "is_chapter_overview": True, "origin": "ai",
        }
        settings = [
            {"chapter_id": "chapter-1", "chapter_number": "1", "chapter_title": "第1章", "overview_policy": "required", "overview_reason": "需要总览"},
            {"chapter_id": "chapter-2", "chapter_number": "2", "chapter_title": "第2章", "overview_policy": "exempt", "overview_reason": "无需图片"},
        ]
        stage3 = {
            "schema_version": 1, "stage": "stage3", "project_id": project_id,
            "stage2_confirmation_sha256": stage2_receipt["confirmation_sha256"],
            "visual_direction": {"palette": "深蓝、红色与白色", "style": "商务", "background": "白色或浅灰底", "density": "适中", "avoid": ["复杂渐变"]}, "chapter_settings": settings, "images": [image], "cleanup_actions": [],
        }
        self._write(confirm_ui.STAGE3_INPUT, stage3)
        confirmed_image = deepcopy(image)
        confirmed_image["name"] = "用户确认后的总览图"
        confirmed_visual = {"palette": "用户确认的蓝红配色", "style": "用户确认的商务风格", "background": "用户确认的浅色背景", "density": "用户确认的适中密度", "avoid": ["用户确认的复杂渐变"]}
        stage3_receipt = self._receipt(confirm_ui.STAGE3_RECEIPT, {
            "schema_version": 1, "stage": "stage3", "status": "confirmed", "project_id": project_id,
            "stage2_confirmation_sha256": stage2_receipt["confirmation_sha256"],
            "source_sha256": confirm_ui.sha256_data(stage3),
            "data": {"visual_direction": confirmed_visual, "chapter_settings": settings, "images": [confirmed_image], "cleanup_actions": []},
            "confirmed_at": confirm_ui.utc_now(),
        })
        delivery = {
            "word_batch_count": 2,
            "word_batches": [
                {"id": "word-batch-1", "order": 1, "chapter_ids": ["chapter-1"], "chapter_numbers": ["1"], "chapter_titles": ["第1章"], "planned_pages": 10, "output_filename": "测试标书-第1批.docx"},
                {"id": "word-batch-2", "order": 2, "chapter_ids": ["chapter-2"], "chapter_numbers": ["2"], "chapter_titles": ["第2章"], "planned_pages": 10, "output_filename": "测试标书-第2批.docx"},
            ],
            "image_plan_workbook": {"count": 1, "format": ".xlsx", "filename": "测试标书-图片规划表.xlsx", "purpose": "交给其他AI生图", "worksheet_names": ["图片规划清单"], "columns": ["图号"], "image_count": 1},
            "skill_boundary": {"generate_word_documents": True, "generate_image_plan_excel": True, "generate_images": False, "insert_images": False},
            "delivery_output_dir": str(self.project_dir),
            "additional_notes": "",
        }
        stage4 = {
            "schema_version": 1, "stage": "stage4", "project_id": project_id, "generated_at": confirm_ui.utc_now(),
            "stage3_confirmation_sha256": stage3_receipt["confirmation_sha256"],
            "summary": {"project_name": "测试标书", "client": "", "project_overview": "", "chapter_count": 2, "planned_pages": 20, "image_count": 1},
            "delivery": delivery,
        }
        self.stage4 = stage4
        self._write(confirm_ui.STAGE4_INPUT, stage4)
        self.stage4_receipt = self._receipt(confirm_ui.STAGE4_RECEIPT, {
            "schema_version": 1, "stage": "stage4", "status": "confirmed", "project_id": project_id,
            "stage3_confirmation_sha256": stage3_receipt["confirmation_sha256"],
            "source_sha256": confirm_ui.sha256_data(stage4), "data": delivery,
            "confirmed_at": confirm_ui.utc_now(),
        })

    def _write_artifacts(self, batch_id: str, marker: str) -> None:
        manifest = protocol.load_manifest(self.project_dir)
        batch = next(item for item in manifest["word_batches"] if item["id"] == batch_id)
        source_path = protocol.delivery_dir(self.project_dir) / batch["source_path"]
        export_path = protocol.delivery_dir(self.project_dir) / batch["export_path"]
        source_path.parent.mkdir(parents=True, exist_ok=True)
        export_path.parent.mkdir(parents=True, exist_ok=True)
        chapter_id = batch["chapter_ids"][0]
        source = {
            "schema_version": 1, "kind": "bid_delivery_source", "project_id": manifest["project_id"],
            "stage4_confirmation_sha256": manifest["stage4_confirmation_sha256"], "batch_id": batch_id,
            "batch_order": batch["order"], "source_version": 1, "generated_at": "2026-08-13T10:00:00+08:00",
            "updated_at": "2026-08-13T10:00:00+08:00", "writing_rules_sha256": manifest["writing_rules"]["project_sha256"],
            "planned_pages": batch["planned_pages"],
            "actual_pages": None, "chapters": [
                {"id": chapter_id, "number": batch["chapter_numbers"][0], "title": batch["chapter_titles"][0], "level": 1, "order": 1}
            ], "blocks": [
                {"id": "block-1", "order": 1, "type": "paragraph", "chapter_id": chapter_id, "text": f"{marker}版正文说明。"},
                {"id": "heading-1", "order": 2, "type": "heading", "chapter_id": chapter_id, "level": 2, "number": f"{batch['chapter_numbers'][0]}.1", "title": "二级标题"},
                {"id": "block-2", "order": 3, "type": "list", "chapter_id": chapter_id, "items": ["要点一", "要点二"]},
                {"id": "block-3", "order": 4, "type": "table", "chapter_id": chapter_id, "columns": ["事项", "说明"], "rows": [["范围", "本批次范围"]]},
                {"id": "density-1", "order": 5, "type": "paragraph", "chapter_id": chapter_id, "text": "本段用于说明项目实施范围、组织机制、工作步骤、质量控制和交付验收要求。" * (4000 if marker == "page-mismatch-fixed" else 1800)},
            ],
        }
        if batch_id == "word-batch-1":
            source["blocks"].append({"id": "block-4", "order": 6, "type": "image_placeholder", "chapter_id": chapter_id, "figure_no": "图1-1", "name": "用户确认后的总览图", "note": "仅为规划占位，不生成图片。"})
        source["blocks"].append({"id": f"block-{len(source['blocks']) + 1}", "order": len(source["blocks"]) + 1, "type": "page_break", "chapter_id": chapter_id})
        source_path.write_text(json.dumps(source, ensure_ascii=False), encoding="utf-8")
        export_path.write_bytes(f"docx-placeholder-{marker}".encode("utf-8"))

    def test_initialization_is_authorized_idempotent_and_empty(self) -> None:
        manifest, created = protocol.initialize_delivery(self.project_dir)
        self.assertTrue(created)
        self.assertEqual(manifest["status"], "preparing")
        self.assertEqual(manifest["word_batch_count"], 2)
        self.assertTrue(all(item["status"] == "pending" for item in manifest["word_batches"]))
        self.assertFalse(any((protocol.delivery_dir(self.project_dir) / name).is_file() for name in ("source/batch-01.json", "exports/测试标书-第1批.docx", "source/image-plan.json")))
        resumed, created_again = protocol.initialize_delivery(self.project_dir)
        self.assertFalse(created_again)
        self.assertEqual(resumed, protocol.load_manifest(self.project_dir))
        self.assertTrue((protocol.delivery_dir(self.project_dir) / protocol.REQUESTS_DIR_NAME).is_dir())
        self.assertTrue((protocol.delivery_dir(self.project_dir) / protocol.RESULTS_DIR_NAME).is_dir())

    def test_stage4_writing_rules_are_snapshotted_and_bound(self) -> None:
        manifest, _created = protocol.initialize_delivery(self.project_dir)
        rules_path = protocol.writing_rules_path(self.project_dir)
        self.assertTrue(rules_path.is_file())
        self.assertEqual(manifest["writing_rules"]["project_sha256"], protocol.sha256_file(rules_path))
        self.assertEqual(manifest["writing_rules"]["source_sha256"], manifest["writing_rules"]["project_sha256"])
        tampered = rules_path.read_bytes() + "\n故意改写规则快照。".encode("utf-8")
        rules_path.write_bytes(tampered)
        with self.assertRaisesRegex(ValueError, "规则缺失或已被替换"):
            protocol.load_manifest(self.project_dir)

    def test_source_must_bind_current_writing_rules(self) -> None:
        protocol.initialize_delivery(self.project_dir)
        protocol.begin_active_batch(self.project_dir)
        self._write_artifacts("word-batch-1", "rules-binding")
        source_path = protocol.delivery_dir(self.project_dir) / "source/batch-01.json"
        source = confirm_ui.read_json(source_path)
        source["writing_rules_sha256"] = "0" * 64
        source_path.write_text(json.dumps(source, ensure_ascii=False), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "未读取当前Stage4标书生成规则"):
            protocol.register_batch_artifacts(self.project_dir, "word-batch-1")

    def test_state_machine_requires_artifacts_and_review(self) -> None:
        protocol.initialize_delivery(self.project_dir)
        with self.assertRaisesRegex(ValueError, "待审校"):
            protocol.confirm_batch(self.project_dir, "word-batch-1")
        protocol.begin_active_batch(self.project_dir)
        with self.assertRaisesRegex(ValueError, "需要先由AI生成"):
            protocol.register_batch_artifacts(self.project_dir, "word-batch-1")
        self._write_artifacts("word-batch-1", "v1")
        manifest = protocol.register_batch_artifacts(self.project_dir, "word-batch-1")
        self.assertEqual(manifest["status"], "awaiting_batch_review")
        manifest, revision = protocol.request_revision(self.project_dir, "word-batch-1", "请补强这一批的实施逻辑")
        self.assertEqual(manifest["status"], "revision_pending")
        self.assertEqual(manifest["pending_request_count"], 1)
        self.assertEqual(revision["type"], "revision")
        waited = protocol.wait_for_event(self.project_dir, "revision", 1, 0)
        self.assertIsNotNone(waited)
        manifest = protocol.begin_revision(self.project_dir, "word-batch-1")
        self.assertEqual(manifest["status"], "revising")
        self.assertEqual(manifest["pending_request_count"], 0)
        self._write_artifacts("word-batch-1", "v2")
        protocol.register_batch_artifacts(self.project_dir, "word-batch-1")
        manifest, confirmed = protocol.confirm_batch(self.project_dir, "word-batch-1")
        self.assertEqual(confirmed["type"], "batch-confirmed")
        self.assertTrue((self.project_dir / "测试标书-第1批.docx").is_file())
        self.assertEqual(manifest["status"], "generating")
        self.assertEqual(manifest["word_batches"][1]["status"], "generating")
        self.assertEqual(manifest["active_batch_id"], "word-batch-2")
        protocol.begin_active_batch(self.project_dir)
        self._write_artifacts("word-batch-2", "v1")
        protocol.register_batch_artifacts(self.project_dir, "word-batch-2")
        manifest, _ = protocol.confirm_batch(self.project_dir, "word-batch-2")
        self.assertEqual(manifest["status"], "all_batches_confirmed")
        self.assertIsNone(manifest["active_batch_id"])

    def test_thin_source_cannot_enter_review(self) -> None:
        protocol.initialize_delivery(self.project_dir)
        protocol.begin_active_batch(self.project_dir)
        self._write_artifacts("word-batch-1", "thin")
        source_path = protocol.delivery_dir(self.project_dir) / "source" / "batch-01.json"
        source = protocol.read_json(source_path)
        next(block for block in source["blocks"] if block["id"] == "density-1")["text"] = "内容偏薄。"
        protocol.atomic_write_json(source_path, source)
        with self.assertRaisesRegex(ValueError, "低于计划10页的允许下限9页"):
            protocol.register_batch_artifacts(self.project_dir, "word-batch-1")

    def test_structure_metrics_reject_excessive_level3_density(self) -> None:
        source = {
            "chapters": [{"id": "chapter-1", "number": "1", "title": "第1章"}],
            "blocks": [
                {"id": "h2", "order": 1, "type": "heading", "chapter_id": "chapter-1", "level": 2, "number": "1.1", "title": "二级标题"},
            ],
        }
        order = 2
        for index in range(1, 9):
            source["blocks"].extend([
                {"id": f"h3-{index}", "order": order, "type": "heading", "chapter_id": "chapter-1", "level": 3, "number": f"1.1.{index}", "title": f"三级标题{index}"},
                {"id": f"p-{index}-1", "order": order + 1, "type": "paragraph", "chapter_id": "chapter-1", "text": "第一段支撑内容。"},
                {"id": f"p-{index}-2", "order": order + 2, "type": "paragraph", "chapter_id": "chapter-1", "text": "第二段支撑内容。"},
            ])
            order += 3
        with self.assertRaisesRegex(ValueError, "三级标题数量超限"):
            protocol.source_structure_metrics(source)

    def test_stage2_outline_lock_rejects_new_level3_heading(self) -> None:
        stage2_path = self.data_dir / confirm_ui.STAGE2_INPUT
        receipt_path = self.data_dir / confirm_ui.STAGE2_RECEIPT
        stage2 = confirm_ui.read_json(stage2_path)
        receipt = confirm_ui.read_json(receipt_path)
        child = {"id": "outline-1-1", "number": "1.1", "title": "已确认二级标题", "level": 2, "order": 1, "children": []}
        stage2["chapters"][0]["children"] = [child]
        receipt["data"]["chapters"][0]["children"] = [child]
        receipt["source_sha256"] = confirm_ui.sha256_data(stage2)
        receipt.pop("confirmation_sha256", None)
        receipt["confirmation_sha256"] = confirm_ui.sha256_data(receipt)
        confirm_ui.atomic_write_json(stage2_path, stage2)
        confirm_ui.atomic_write_json(receipt_path, receipt)
        source = {
            "chapters": [{"id": "chapter-1", "number": "1", "title": "第1章"}],
            "blocks": [
                {"type": "heading", "chapter_id": "chapter-1", "level": 2, "number": "1.1", "title": "已确认二级标题"},
                {"type": "heading", "chapter_id": "chapter-1", "level": 3, "number": "1.1.1", "title": "正文阶段新增三级标题"},
            ],
        }
        batch = {"chapter_ids": ["chapter-1"]}
        with self.assertRaisesRegex(ValueError, "阶段2重新确认目录"):
            protocol.validate_source_against_confirmed_outline(self.project_dir, source, batch)

    def test_source_change_invalidates_existing_manifest(self) -> None:
        protocol.initialize_delivery(self.project_dir)
        changed = deepcopy(self.stage4)
        changed["generated_at"] = "2099-01-01T00:00:00+08:00"
        self._write(confirm_ui.STAGE4_INPUT, changed)
        with self.assertRaisesRegex(ValueError, "最终交付授权"):
            protocol.load_manifest(self.project_dir)

    def test_local_word_and_single_sheet_image_plan_exports(self) -> None:
        """Exercise concrete deliverables without generating any image."""
        protocol.initialize_delivery(self.project_dir)
        protocol.begin_active_batch(self.project_dir)
        self._write_artifacts("word-batch-1", "export")
        word_result = export_word.export_word(self.project_dir, "word-batch-1")
        word_path = Path(word_result["output"])
        self.assertTrue(word_path.is_file())
        self.assertTrue(zipfile.is_zipfile(word_path))
        self.assertEqual(word_result["manifest"]["word_batches"][0]["status"], "ready_for_review")
        from docx import Document
        document = Document(word_path)
        self.assertTrue(any(item.text == "1.1 二级标题" and item.style.name == "Heading 2" for item in document.paragraphs))
        image_result = export_image_plan.export_image_plan(self.project_dir, render_preview=True)
        excel_path = Path(image_result["output"])
        self.assertTrue(excel_path.is_file())
        self.assertTrue(zipfile.is_zipfile(excel_path))
        self.assertEqual(image_result["validation"]["checks"]["worksheet_count"], 1)
        self.assertFalse(image_result["validation"]["checks"]["image_generation_in_scope"])
        image_source = protocol.read_json(protocol.delivery_dir(self.project_dir) / "source" / "image-plan.json")
        self.assertEqual(image_source["images"][0]["name"], "用户确认后的总览图")
        self.assertEqual(image_source["visual_direction"]["style"], "用户确认的商务风格")
        self.assertTrue(image_source["images"][0]["ai_prompt"])
        from openpyxl import load_workbook
        image_workbook = load_workbook(excel_path, read_only=True, data_only=True)
        self.assertEqual(image_workbook.active["O4"].value, "AI生图提示词")
        self.assertTrue(image_workbook.active["O5"].value)

    def test_final_delivery_locks_exact_artifacts_without_wps_receipts(self) -> None:
        protocol.initialize_delivery(self.project_dir)
        export_image_plan.export_image_plan(self.project_dir)
        manifest, _event = protocol.confirm_image_plan(self.project_dir)
        self.assertEqual(manifest["image_plan_workbook"]["status"], "confirmed")
        for batch_id in ("word-batch-1", "word-batch-2"):
            protocol.begin_active_batch(self.project_dir)
            self._write_artifacts(batch_id, f"final-{batch_id}")
            export_word.export_word(self.project_dir, batch_id)
            with self.assertRaisesRegex(ValueError, "当前交付尚未|最终交付条件"):
                protocol.confirm_final_delivery(self.project_dir)
            manifest, _event = protocol.confirm_batch(self.project_dir, batch_id)
        self.assertEqual(manifest["status"], "all_batches_confirmed")
        final = protocol.final_delivery_payload(self.project_dir)
        self.assertFalse(final["eligible"])
        protocol.record_wps_page_check(self.project_dir, "word-batch-1", 10, "测试WPS")
        protocol.record_wps_page_check(self.project_dir, "word-batch-2", 10, "测试WPS")
        final = protocol.final_delivery_payload(self.project_dir)
        self.assertTrue(final["eligible"])
        confirmed, event = protocol.confirm_final_delivery(self.project_dir)
        self.assertEqual(confirmed["status"], "final_confirmed")
        self.assertEqual(event["type"], "final-confirmed")
        self.assertTrue((protocol.delivery_dir(self.project_dir) / "confirmations" / "final-confirmation.json").is_file())
        self.assertTrue((self.project_dir / "测试标书-图片规划表.xlsx").is_file())
        with self.assertRaisesRegex(ValueError, "最终交付已确认"):
            protocol.apply_direct_edit(self.project_dir, "word-batch-1", "block-1", confirmed["word_batches"][0]["source_sha256"], "不应写入的最终后修改。")

    def test_in_range_ai_estimate_allows_final_delivery_without_manual_wps_entry(self) -> None:
        protocol.initialize_delivery(self.project_dir)
        export_image_plan.export_image_plan(self.project_dir)
        protocol.confirm_image_plan(self.project_dir)
        for batch_id in ("word-batch-1", "word-batch-2"):
            protocol.begin_active_batch(self.project_dir)
            self._write_artifacts(batch_id, "estimate-only")
            export_word.export_word(self.project_dir, batch_id)
            manifest = protocol.load_manifest(self.project_dir)
            batch = next(item for item in manifest["word_batches"] if item["id"] == batch_id)
            validation_path = protocol.delivery_dir(self.project_dir) / protocol.RESULTS_DIR_NAME / f"word-batch-{batch['order']:02d}-validation.json"
            validation = protocol.read_json(validation_path)
            validation["page_verification"]["estimated_pages"] = batch["planned_pages"]
            protocol.atomic_write_json(validation_path, validation)
            protocol.confirm_batch(self.project_dir, batch_id)
        final = protocol.final_delivery_payload(self.project_dir)
        self.assertTrue(final["eligible"], final)

    def test_reader_uses_valid_source_and_paginates_blocks(self) -> None:
        protocol.initialize_delivery(self.project_dir)
        protocol.begin_active_batch(self.project_dir)
        self._write_artifacts("word-batch-1", "reader")
        protocol.register_batch_artifacts(self.project_dir, "word-batch-1")
        overview = protocol.delivery_overview_payload(self.project_dir)
        self.assertTrue(overview["read_only"])
        self.assertTrue(overview["batches"][0]["readable"])
        self.assertFalse(overview["batches"][1]["readable"])
        page1 = protocol.batch_reader_payload(self.project_dir, "word-batch-1", offset=0, limit=2)
        self.assertEqual(len(page1["blocks"]), 2)
        self.assertEqual(page1["paging"]["next_offset"], 2)
        page2 = protocol.batch_reader_payload(self.project_dir, "word-batch-1", offset=2, limit=80)
        self.assertEqual(len(page2["blocks"]), 5)
        self.assertIsNone(page2["paging"]["next_offset"])
        with self.assertRaisesRegex(ValueError, "最多读取"):
            protocol.batch_reader_payload(self.project_dir, "word-batch-1", limit=81)

    def test_direct_edit_creates_history_and_invalidates_word_export(self) -> None:
        protocol.initialize_delivery(self.project_dir)
        protocol.begin_active_batch(self.project_dir)
        self._write_artifacts("word-batch-1", "direct")
        protocol.register_batch_artifacts(self.project_dir, "word-batch-1")
        before = protocol.load_manifest(self.project_dir)
        result = protocol.apply_direct_edit(self.project_dir, "word-batch-1", "block-1", before["word_batches"][0]["source_sha256"], "经修订后的确定性正文。")
        updated = result["manifest"]
        batch = updated["word_batches"][0]
        self.assertEqual(updated["status"], "export_pending")
        self.assertEqual(batch["status"], "export_pending")
        self.assertIsNone(batch["export_sha256"])
        self.assertTrue((protocol.delivery_dir(self.project_dir) / "history" / f"{result['record_id']}.json").is_file())
        reader = protocol.batch_reader_payload(self.project_dir, "word-batch-1")
        self.assertEqual(reader["blocks"][0]["text"], "经修订后的确定性正文。")
        with self.assertRaisesRegex(ValueError, "需重新导出"):
            protocol.apply_direct_edit(self.project_dir, "word-batch-1", "block-1", before["word_batches"][0]["source_sha256"], "旧页面覆盖。")

    def test_direct_edit_supports_list_and_table(self) -> None:
        protocol.initialize_delivery(self.project_dir)
        protocol.begin_active_batch(self.project_dir)
        self._write_artifacts("word-batch-1", "rich-direct")
        protocol.register_batch_artifacts(self.project_dir, "word-batch-1")
        manifest = protocol.load_manifest(self.project_dir)
        result = protocol.apply_direct_edit(self.project_dir, "word-batch-1", "block-2", manifest["word_batches"][0]["source_sha256"], replacement_items=["新要点一", "新要点二"])
        self.assertEqual(result["manifest"]["word_batches"][0]["status"], "export_pending")
        export_word.export_word(self.project_dir, "word-batch-1")
        manifest = protocol.load_manifest(self.project_dir)
        result = protocol.apply_direct_edit(self.project_dir, "word-batch-1", "block-3", manifest["word_batches"][0]["source_sha256"], replacement_columns=["项目", "要求"], replacement_rows=[["范围", "全量"]])
        self.assertEqual(result["manifest"]["word_batches"][0]["status"], "export_pending")

    def test_editing_a_confirmed_batch_archives_its_receipt_and_reopens_review(self) -> None:
        protocol.initialize_delivery(self.project_dir)
        protocol.begin_active_batch(self.project_dir)
        self._write_artifacts("word-batch-1", "reopen")
        protocol.register_batch_artifacts(self.project_dir, "word-batch-1")
        confirmed, _event = protocol.confirm_batch(self.project_dir, "word-batch-1")
        updated = protocol.apply_direct_edit(self.project_dir, "word-batch-1", "block-1", confirmed["word_batches"][0]["source_sha256"], "确认后的修订正文。")
        self.assertEqual(updated["manifest"]["status"], "export_pending")
        self.assertEqual(updated["manifest"]["active_batch_id"], "word-batch-1")
        self.assertFalse((protocol.delivery_dir(self.project_dir) / "confirmations" / "batch-01-confirmation.json").exists())
        self.assertTrue(list((protocol.delivery_dir(self.project_dir) / "history").glob("batch-01-confirmation-invalidated-*.json")))

    def test_wps_page_mismatch_reopens_delivery_as_review_version(self) -> None:
        protocol.initialize_delivery(self.project_dir)
        protocol.begin_active_batch(self.project_dir)
        self._write_artifacts("word-batch-1", "page-mismatch")
        export_word.export_word(self.project_dir, "word-batch-1")
        confirmed, _event = protocol.confirm_batch(self.project_dir, "word-batch-1")
        self.assertEqual(confirmed["word_batches"][0]["status"], "confirmed")
        result = protocol.record_wps_page_check(self.project_dir, "word-batch-1", 5, "测试WPS")
        self.assertTrue(result["revision_required"])
        self.assertEqual(result["revision_action"], "expand")
        self.assertTrue(result["replan_required"])
        calibration = protocol.read_json(protocol.page_calibration_path(self.project_dir))
        self.assertEqual(calibration["sample_count"], 1)
        self.assertLess(calibration["ratio"], 1)
        self.assertEqual(result["manifest"]["status"], "revision_pending")
        batch = result["manifest"]["word_batches"][0]
        self.assertEqual(batch["status"], "revision_pending")
        self.assertEqual(batch["output_filename"], "测试标书-第1批—审校版1.docx")
        self.assertEqual(result["event"]["payload"]["kind"], "page-count-mismatch")
        self.assertFalse((protocol.delivery_dir(self.project_dir) / "confirmations" / "batch-01-confirmation.json").exists())
        self.assertTrue(list((protocol.delivery_dir(self.project_dir) / "history").glob("batch-01-page-mismatch-*.json")))

        protocol.begin_revision(self.project_dir, "word-batch-1")
        self._write_artifacts("word-batch-1", "page-mismatch-fixed")
        revised = protocol.register_batch_artifacts(self.project_dir, "word-batch-1")
        self.assertEqual(revised["word_batches"][0]["status"], "ready_for_review")
        self.assertEqual(revised["word_batches"][0]["output_filename"], "测试标书-第1批—审校版1.docx")

    def test_ai_request_lifecycle_rejects_stale_source(self) -> None:
        protocol.initialize_delivery(self.project_dir)
        protocol.begin_active_batch(self.project_dir)
        self._write_artifacts("word-batch-1", "ai")
        protocol.register_batch_artifacts(self.project_dir, "word-batch-1")
        manifest = protocol.load_manifest(self.project_dir)
        created = protocol.create_ai_request(self.project_dir, "word-batch-1", "block-1", manifest["word_batches"][0]["source_sha256"], "请将本段改得更严谨。")
        request_id = created["request"]["id"]
        self.assertEqual(created["manifest"]["status"], "revision_pending")
        self.assertEqual(len(protocol.list_ai_requests(self.project_dir, "word-batch-1")), 1)
        started = protocol.begin_ai_request(self.project_dir, request_id)
        self.assertEqual(started["manifest"]["status"], "revising")
        applied = protocol.apply_ai_request_result(self.project_dir, request_id, "经AI处理后的严谨正文。")
        self.assertEqual(applied["manifest"]["status"], "export_pending")
        self.assertTrue((protocol.delivery_dir(self.project_dir) / "results" / f"{request_id}-result.json").is_file())
        self.assertEqual(protocol.list_ai_requests(self.project_dir)[0]["status"], "applied")

    def test_stale_ai_request_is_superseded_not_applied(self) -> None:
        protocol.initialize_delivery(self.project_dir)
        protocol.begin_active_batch(self.project_dir)
        self._write_artifacts("word-batch-1", "stale")
        protocol.register_batch_artifacts(self.project_dir, "word-batch-1")
        manifest = protocol.load_manifest(self.project_dir)
        created = protocol.create_ai_request(self.project_dir, "word-batch-1", "block-1", manifest["word_batches"][0]["source_sha256"], "请修改本段。")
        request_id = created["request"]["id"]
        request_path = protocol.delivery_dir(self.project_dir) / "requests" / f"{request_id}.json"
        stale_request = protocol.read_json(request_path)
        stale_request["source_sha256"] = "0" * 64
        protocol.atomic_write_json(request_path, stale_request)
        with self.assertRaisesRegex(ValueError, "已自动失效"):
            protocol.begin_ai_request(self.project_dir, request_id)
        self.assertEqual(protocol.read_json(request_path)["status"], "superseded")
        self.assertEqual(protocol.load_manifest(self.project_dir)["status"], "export_pending")

    def test_image_plan_can_be_edited_and_reexported_for_review(self) -> None:
        protocol.initialize_delivery(self.project_dir)
        export_image_plan.export_image_plan(self.project_dir)
        before = protocol.image_plan_payload(self.project_dir)
        self.assertFalse(before["read_only"])
        result = protocol.apply_image_plan_direct_edit(
            self.project_dir, "image-1", before["source_sha256"],
            {"name": "更新后的总览图", "type": "流程图", "purpose": "说明实施闭环", "core_nodes": ["目标", "闭环"],
             "composition": "横向流程", "orientation": "landscape", "is_chapter_overview": True, "placement_note": "第1章导语后"},
        )
        self.assertEqual(result["manifest"]["image_plan_workbook"]["status"], "export_pending")
        changed = protocol.image_plan_payload(self.project_dir)
        self.assertEqual(changed["images"][0]["name"], "更新后的总览图")
        export_image_plan.export_image_plan(self.project_dir)
        self.assertEqual(protocol.image_plan_payload(self.project_dir)["status"], "ready_for_review")

    def test_image_plan_ai_request_lifecycle(self) -> None:
        protocol.initialize_delivery(self.project_dir)
        export_image_plan.export_image_plan(self.project_dir)
        before = protocol.image_plan_payload(self.project_dir)
        created = protocol.create_image_plan_ai_request(self.project_dir, "image-1", before["source_sha256"], "请突出闭环与验收。")
        request_id = created["request"]["id"]
        started = protocol.begin_ai_request(self.project_dir, request_id)
        self.assertEqual(started["request"]["kind"], "bid_delivery_image_plan_ai_request")
        result = protocol.apply_image_plan_ai_request_result(self.project_dir, request_id, {"name": "AI改写后的总览图", "type": "流程图", "purpose": "展示服务闭环", "core_nodes": ["目标", "闭环"], "composition": "横向流程", "orientation": "landscape", "is_chapter_overview": True, "placement_note": "第1章导语后"})
        self.assertEqual(result["manifest"]["image_plan_workbook"]["status"], "export_pending")
        export_image_plan.export_image_plan(self.project_dir)
        self.assertEqual(protocol.image_plan_payload(self.project_dir)["images"][0]["name"], "AI改写后的总览图")

    def test_confirmed_image_plan_is_published_and_reopens_on_edit(self) -> None:
        protocol.initialize_delivery(self.project_dir)
        export_image_plan.export_image_plan(self.project_dir)
        manifest, event = protocol.confirm_image_plan(self.project_dir)
        workbook = manifest["image_plan_workbook"]
        self.assertEqual(event["type"], "image-plan-confirmed")
        self.assertEqual(workbook["status"], "confirmed")
        self.assertTrue((self.project_dir / workbook["filename"]).is_file())
        payload = protocol.image_plan_payload(self.project_dir)
        self.assertTrue(payload["read_only"])
        changed = protocol.apply_image_plan_direct_edit(
            self.project_dir, "image-1", payload["source_sha256"],
            {"name": "重新打开后的总览图", "type": "流程图", "purpose": "说明实施闭环", "core_nodes": ["目标", "闭环"],
             "composition": "横向流程", "orientation": "landscape", "is_chapter_overview": True, "placement_note": "第1章导语后"},
        )
        self.assertEqual(changed["manifest"]["image_plan_workbook"]["status"], "export_pending")
        self.assertFalse((protocol.delivery_dir(self.project_dir) / "confirmations" / "image-plan-confirmation.json").exists())

    def test_manifest_rejects_forged_completion_and_path_escape(self) -> None:
        manifest, _ = protocol.initialize_delivery(self.project_dir)
        forged = deepcopy(manifest)
        forged["status"] = "final_confirmed"
        forged["active_batch_id"] = None
        with self.assertRaisesRegex(ValueError, "最终阶段"):
            protocol.validate_manifest(forged, self.project_dir)
        escaped = deepcopy(manifest)
        escaped["word_batches"][0]["source_path"] = "../outside.json"
        with self.assertRaisesRegex(ValueError, "Word源稿路径"):
            protocol.validate_manifest(escaped, self.project_dir)


if __name__ == "__main__":
    unittest.main()
